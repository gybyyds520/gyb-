# -*- coding: utf-8 -*-
"""试卷文本提取通道：PDF(含扫描)/Word/图片 -> 纯文本 JSON。
用法: python paper_extract.py <input_path> <output_json_path>
输出: {"ok": true|false, "text": "...", "pages": N, "method": "PDF_TEXT|PDF_OCR|DOCX|IMAGE_OCR", "error": "..."}
"""
import json
import os
import sys


def extract_pdf_text(path):
    import fitz
    doc = fitz.open(path)
    pages_text = []
    for page in doc:
        pages_text.append(page.get_text("text"))
    total = "".join(pages_text).strip()
    if len(total) < 30 and doc.page_count > 0:
        # 判定为扫描版 PDF：逐页渲染 -> OCR
        imgs = []
        for page in doc:
            pix = page.get_pixmap(dpi=200)
            tmp = path + f".p{page.number}.png"
            pix.save(tmp)
            imgs.append(tmp)
        text = ocr_images(imgs)
        for t in imgs:
            try:
                os.remove(t)
            except OSError:
                pass
        return text, doc.page_count, "PDF_OCR"
    return total, doc.page_count, "PDF_TEXT"


def extract_docx_text(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for tbl in d.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join([c for c in cells if c]))
    return "\n".join(parts), 1, "DOCX"


_ocr = None


def ocr_images(img_paths):
    global _ocr
    if _ocr is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr = RapidOCR()
    lines = []
    for img in img_paths:
        out = _ocr(img)
        result = out[0] if isinstance(out, tuple) else out
        if not result:
            continue
        for row in result:
            if row and len(row) >= 2:
                text = str(row[1]).strip()
                if text:
                    lines.append(text)
    return "\n".join(lines)


def main():
    if len(sys.argv) < 3:
        print(json.dumps({"ok": False, "error": "usage: paper_extract.py <input> <output>"}, ensure_ascii=False))
        return 1
    src = sys.argv[1]
    out = sys.argv[2]
    if not os.path.exists(src):
        print(json.dumps({"ok": False, "error": "file not found: " + src}, ensure_ascii=False))
        return 1
    ext = os.path.splitext(src)[1].lower()
    try:
        if ext == ".pdf":
            text, pages, method = extract_pdf_text(src)
        elif ext in (".docx", ".doc"):
            text, pages, method = extract_docx_text(src)
        elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".webp"):
            text = ocr_images([src])
            pages, method = 1, "IMAGE_OCR"
        else:
            print(json.dumps({"ok": False, "error": "unsupported type: " + ext}, ensure_ascii=False))
            return 1
        result = {"ok": True, "text": text or "", "pages": pages, "method": method}
        with open(out, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False)
        return 0
    except Exception as e:  # noqa
        result = {"ok": False, "error": str(e)}
        with open(out, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False)
        return 2


if __name__ == "__main__":
    sys.exit(main())
