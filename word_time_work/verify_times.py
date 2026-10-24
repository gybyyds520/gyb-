# -*- coding: utf-8 -*-
"""验证最终时间：输出所有含日期的单元格/段落文本"""
from docx import Document

BASE = r'C:\Users\郭耀彬\Desktop\gybct\gybxx'
for path, label in [
    (BASE + r'\2301132013-郭耀彬-题目审核表.docx', '题目审核表'),
    (BASE + r'\2301132013-郭耀彬-开题报告.docx', '开题报告'),
]:
    doc = Document(path)
    print(f'===== {label} =====')
    seen = set()
    for p in doc.paragraphs:
        if '202' in p.text and p.text.strip() not in seen:
            seen.add(p.text.strip())
            print('P:', p.text.strip()[:80])
    for ti, table in enumerate(doc.tables):
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if '202' in t and t not in seen:
                    seen.add(t)
                    print(f'T{ti}:', t[:100].replace('\n', ' | '))
