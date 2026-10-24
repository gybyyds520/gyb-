# -*- coding: utf-8 -*-
"""检查两份 DOCX 中时间文本的 run 结构，确保修改时保留原格式"""
import sys
from docx import Document

BASE = r'C:\Users\郭耀彬\Desktop\gybct\gybxx'

def scan(path, targets):
    doc = Document(path)
    found = []
    # 遍历正文段落与表格单元格
    def check_text(text, where):
        for t in targets:
            if t in text:
                found.append((where, text.strip()[:120]))
    for i, p in enumerate(doc.paragraphs):
        check_text(p.text, f'P{i}')
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                # 单元格内段落
                for pi, p in enumerate(cell.paragraphs):
                    check_text(p.text, f'T{ti}r{ri}c{ci}p{pi}')
                    # run 结构
                    if any(t in p.text for t in targets):
                        runs = [(r.text, r.bold, r.font.size.pt if r.font.size else None) for r in p.runs]
                        print(f'  RUNS[{ti}r{ri}c{ci}p{pi}]: {runs}')
    for w in found:
        print(w)

print('===== 题目审核表 =====')
scan(BASE + r'\2301132013-郭耀彬-题目审核表.docx', ['2026年8月13日', '2026年 8月15日'])
print('===== 开题报告 =====')
scan(BASE + r'\2301132013-郭耀彬-开题报告.docx', ['2026年10月11日', '2026年 10月 11日'])
